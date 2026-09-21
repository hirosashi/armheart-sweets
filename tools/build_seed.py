#!/usr/bin/env python3
"""エンドユーザー提供のExcel（配合表・原価計算表・発注書）から db/seed_real.sql を生成する。

使い方:  python3 tools/build_seed.py
入力:    tools/source/ に置いた4ファイル（下の SRC を参照）
出力:    db/seed_real.sql
"""
import datetime
import re
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
SRC = Path(__file__).resolve().parent / 'source'
RECIPE = SRC / '【改訂】最終配合表（クロミ_クリスマスチョコケーキ）バッチ配合量追加.xlsx'
COST_AH = SRC / '00.原価計算表アルムハート.xlsx'
COST_IC = SRC / '00.原価計算表伊藤忠_クロミ⑤.xlsx'
ORDER = SRC / 'お餅シート発注書.xlsx'
OUT = BASE / 'db' / 'seed_real.sql'

NO_STOCK = {'水'}
MAKER_PREFIX = re.compile(r'^(販売者|輸入者|製造者)：')


def q(v):
    if v is None or v == '':
        return 'NULL'
    return "'" + str(v).replace('\\', '\\\\').replace("'", "''") + "'"


def num(v):
    return 'NULL' if v is None else repr(float(v))


def clean(v):
    if v is None:
        return None
    s = str(v).replace('\u3000', ' ').strip()
    return s or None


class Registry:
    """材料と業者を名前で一元管理する。"""

    def __init__(self):
        self.materials = {}   # name -> dict
        self.suppliers = {}   # name -> dict

    def supplier(self, raw, stype='purchase'):
        name = clean(raw)
        if not name or name == 'ー' or name == '0' or name == 'メーカー未定':
            return None
        first = MAKER_PREFIX.sub('', name.split('/')[0]).strip()
        if not first:
            return None
        s = self.suppliers.setdefault(first, {'name': first, 'type': stype})
        if s['type'] != stype:
            s['type'] = 'both'
        return first

    def material(self, raw_name, maker=None, allergens=None, kind='material',
                 unit='g', is_supplied=0,
                 purchase_unit=None, purchase_qty=None):
        name = clean(raw_name)
        if not name:
            return None
        parts = [p.strip() for p in name.split('/') if p.strip()]
        key = parts[0]
        alias = '/'.join(parts[1:]) or None
        m = self.materials.get(key)
        if m is None:
            m = {
                'name': key, 'alias': alias, 'kind': kind, 'unit': unit,
                'maker': clean(maker), 'allergens': clean(allergens),
                'supplier': self.supplier(maker), 'is_supplied': is_supplied,
                'stock': 0 if key in NO_STOCK else 1,
                'purchase_unit': purchase_unit, 'purchase_qty': purchase_qty,
            }
            self.materials[key] = m
        else:
            if alias and not m['alias']:
                m['alias'] = alias
            for field, value in (('maker', clean(maker)), ('allergens', clean(allergens)),
                                 ('purchase_unit', purchase_unit)):
                if value and not m[field]:
                    m[field] = value
            if purchase_qty is not None and m['purchase_qty'] is None:
                m['purchase_qty'] = purchase_qty
            if m['supplier'] is None:
                m['supplier'] = self.supplier(maker)
            if kind != 'material':
                m['kind'] = kind
        return key


reg = Registry()

# ---------------------------------------------------------------- 原材料マスタ
# 原価計算表は材料名・メーカーの取得にのみ使う（単価は取り込まない）
wb = openpyxl.load_workbook(COST_AH, data_only=True)
ws = wb['マスタ(アルムハート)']
for row in ws.iter_rows(min_row=2, values_only=True):
    name, maker = row[1], row[2]
    if not clean(name):
        continue
    reg.material(name, maker=maker, purchase_unit='kg', purchase_qty=1000)

# アレルゲンは配合表のマスタが持っている
wb_r = openpyxl.load_workbook(RECIPE, data_only=True)
ws = wb_r['マスタ']
for row in ws.iter_rows(min_row=2, values_only=True):
    name, maker, allergen = row[1], row[2], row[3]
    if not clean(name):
        continue
    reg.material(name, maker=maker, allergens=allergen, purchase_unit='kg', purchase_qty=1000)

# ------------------------------------------------------------------ 部位（パーツ）
parts = []       # {name, batch_total, shared, materials:[(material, qty)]}


def read_recipe_block(ws, first_row, last_row, name_col='B', mat_col='D',
                      qty_col='E', maker_col='G', allergen_col='H', shared=0):
    """部位名が入った行から次の部位までを1バッチの配合として読む。"""
    current = None
    for r in range(first_row, last_row + 1):
        pname = clean(ws[f'{name_col}{r}'].value)
        mat = clean(ws[f'{mat_col}{r}'].value)
        qty = ws[f'{qty_col}{r}'].value
        if pname and mat:
            current = {'name': pname, 'batch_total': 0.0, 'shared': shared, 'materials': []}
            parts.append(current)
        if current is None or not mat:
            continue
        if not isinstance(qty, (int, float)):
            continue  # 「１枚/台」など数値でないものは商品直接配合として別途登録
        key = reg.material(mat, maker=ws[f'{maker_col}{r}'].value,
                           allergens=ws[f'{allergen_col}{r}'].value if allergen_col else None,
                           purchase_unit='kg', purchase_qty=1000)
        current['materials'].append((key, float(qty)))
        current['batch_total'] += float(qty)


ws = wb_r['一般']
read_recipe_block(ws, 16, 57)

ws = wb_r['スポンジ']
read_recipe_block(ws, 14, 45, maker_col='F', allergen_col='G', shared=1)

# クレープは右側の表（I〜K列）に34g充填/枚の配合が入っている
ws = wb_r['クレープ']
crepe = {'name': 'クレープ（34g充填/枚）', 'batch_total': 0.0, 'shared': 1, 'materials': []}
for r in range(21, 31):
    mat = clean(ws[f'I{r}'].value)
    qty = ws[f'K{r}'].value
    if not mat or not isinstance(qty, (int, float)):
        continue
    key = reg.material(mat, maker=ws[f'J{r}'].value, purchase_unit='kg', purchase_qty=1000)
    crepe['materials'].append((key, float(qty)))
    crepe['batch_total'] += float(qty)
parts.append(crepe)

# 「１枚/台」のように数量が数値でない部位（クロミ顔チョコ）は商品直接配合として扱う
parts[:] = [p for p in parts if p['materials']]

# ---------------------------------------------------- 商品への直接配合（枚数もの・資材）
direct = []   # (material, qty, unit)
ws = wb_r['一般']
for r in (53, 54):   # クロミフェイス / クロミドクロ … 1枚/台
    key = reg.material(ws[f'D{r}'].value, maker=ws[f'G{r}'].value,
                       allergens=ws[f'H{r}'].value, unit='枚',
                       purchase_unit='枚', purchase_qty=1)
    direct.append((key, 1.0))

for r in range(60, 64):    # 生産用資材
    name = clean(ws[f'C{r}'].value)
    if not name:
        continue
    key = reg.material(name, kind='pack_production', unit='枚',
                       maker=ws[f'G{r}'].value, purchase_unit='枚', purchase_qty=1)
    direct.append((key, float(ws[f'E{r}'].value or 1)))

for r in range(67, 73):    # 商品用資材（支給品を含む）
    name = clean(ws[f'C{r}'].value)
    qty = ws[f'E{r}'].value
    if not name or not isinstance(qty, (int, float)):
        continue
    supplied = 1 if clean(ws[f'F{r}'].value) == '支給' else 0
    key = reg.material(name, kind='pack_product', unit='枚', maker=ws[f'G{r}'].value,
                       is_supplied=supplied, purchase_unit='枚', purchase_qty=1)
    reg.materials[key]['is_supplied'] = supplied
    if clean(ws[f'G{r}'].value):
        reg.supplier(ws[f'G{r}'].value, 'sales' if supplied else 'purchase')
    direct.append((key, float(qty)))

# ------------------------------------------------- 商品→部位（充填量・取り数）
# 充填量と取り数は原価計算表（伊藤忠_クロミ⑤）に入っている
PRODUCT_PARTS = [
    ('４号茶スポンジ', 150, 2, 1, '1.5cm厚2枚取り1枚使用'),
    ('φ7cm茶スポンジ', 2000, 56, 1, 'φ7 1スライス28個取り／合計56個取り/台'),
    ('トッピング', 6, 1, 1, '配合表は充填量6gに対し配合量4g。差異があるため要確認'),
    ('グラサージュ', 85, 1, 1, None),
    ('チョコレートムース', 160, 1, 1, None),
    ('滑り止め', 3, 1, 1, '接着用'),
]

# お餅シート（発注書の実例）
reg.material('お餅シートプレーンS-5.5-5.5', maker='株式会社ホーライ', kind='material',
             unit='ケース', purchase_unit='ケース', purchase_qty=1)
reg.supplier('株式会社ホーライ')
reg.supplier('伊藤忠食品', 'sales')

# ------------------------------------------------------------------------ 出力
sup_ids = {}
for i, s in enumerate(sorted(reg.suppliers.values(), key=lambda x: x['name']), 1):
    sup_ids[s['name']] = i
mat_ids = {}
for i, key in enumerate(reg.materials.keys(), 1):
    mat_ids[key] = i
part_ids = {}
for i, p in enumerate(parts, 1):
    part_ids[p['name']] = i
PRODUCT_ID = 1

out = ["-- 実データ初期投入（エンドユーザー提供のExcelから tools/build_seed.py で生成）",
       "-- 生成日時: " + datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
       "-- ※ 名前の重複を避けるため、業者・原材料・部位・商品はIDを明示して登録する。",
       "SET NAMES utf8mb4;",
       "SET time_zone = '+09:00';",
       ""]

out.append("-- 自社（発注書の発注者欄）")
out.append("INSERT INTO companies (id, name, address, tel, fax, delivery_place, is_default, sort_no) VALUES")
out.append("  (1, '株式会社アルムハート', '栃木県下都賀郡壬生町大字安塚3343-11', '0282-25-8085', '0282-25-7768', '株式会社アルムハート', 1, 1),")
out.append("  (2, '株式会社エスワイプラス', '栃木県宇都宮市東宿郷4-2-7アークビル9階', '028-678-6058', '028-678-6248', '株式会社エスワイプラス　宇都宮工房', 0, 2);")
out.append("")

out.append("-- 業者（メーカー・発注先・販売先）")
rows = [f"  ({sup_ids[s['name']]}, {q(s['name'])}, {q(s['type'])})"
        for s in sorted(reg.suppliers.values(), key=lambda x: x['name'])]
out.append("INSERT INTO suppliers (id, name, supplier_type) VALUES")
out.append(",\n".join(rows) + ";")
out.append("")

out.append("-- 原材料・資材")
rows = []
for key, m in reg.materials.items():
    rows.append("  (" + ", ".join([
        str(mat_ids[key]), q(m['name']), q(m['alias']), q(m['kind']), q(m['maker']),
        q(m['allergens']), q(m['unit']), q(m['purchase_unit']), num(m['purchase_qty']),
        str(m['is_supplied']), str(m['stock']),
        str(sup_ids[m['supplier']]) if m['supplier'] in sup_ids else 'NULL',
    ]) + ")")
out.append("INSERT INTO materials (id, name, alias_names, kind, maker_name, allergens, unit,"
           " purchase_unit, purchase_qty,"
           " is_supplied, is_stock_managed, supplier_id) VALUES")
out.append(",\n".join(rows) + ";")
out.append("")

out.append("-- 部位（パーツ）＝1バッチの配合")
rows = [f"  ({part_ids[p['name']]}, {q(p['name'])}, 'g', {round(p['batch_total'], 3)!r}, 0.900, 1, {p['shared']})"
        for p in parts]
out.append("INSERT INTO parts (id, name, unit, batch_total_qty, yield_rate, round_batch, is_shared) VALUES")
out.append(",\n".join(rows) + ";")
out.append("")

out.append("-- 部位→原材料（バッチ配合量）")
rows = []
for p in parts:
    for i, (mat, qty) in enumerate(p['materials'], 1):
        rows.append(f"  ({part_ids[p['name']]}, {mat_ids[mat]}, {qty!r}, {i})")
out.append("INSERT INTO part_materials (part_id, material_id, qty, sort_no) VALUES")
out.append(",\n".join(rows) + ";")
out.append("")

out.append("-- 商品")
out.append("INSERT INTO products (id, name, seller_name, spec, shelf_life, launch_date, plan_qty_note,"
           " allergens, author, revised_at, note) VALUES")
out.append("  (1, 'クロミ　クリスマスチョコケーキ', '伊藤忠食品', '４号ホール', '2027/04/03（配合表の値）',"
           " '2026-12-04', '3000～8000台', '小麦・卵・乳成分・大豆・ゼラチン', '河野', '2027-07-24',"
           " '入り数は未定。ガムテープ色＝冷チルは色付き／冷凍は茶色。');")
out.append("")

out.append("-- 商品→部位（充填量・取り数）")
rows = []
for i, (pname, fill, per_fill, use, note) in enumerate(PRODUCT_PARTS, 1):
    rows.append(f"  ({PRODUCT_ID}, {part_ids[pname]}, {float(fill)!r}, 'g', {per_fill}, {use}, {q(note)}, {i})")
out.append("INSERT INTO product_parts (product_id, part_id, fill_qty, fill_unit,"
           " pieces_per_fill, use_pieces, note, sort_no) VALUES")
out.append(",\n".join(rows) + ";")
out.append("")

out.append("-- 商品→原材料・資材の直接使用（1台あたり）")
rows = [f"  ({PRODUCT_ID}, {mat_ids[mat]}, {qty!r}, {i})" for i, (mat, qty) in enumerate(direct, 1)]
out.append("INSERT INTO product_materials (product_id, material_id, qty, sort_no) VALUES")
out.append(",\n".join(rows) + ";")
out.append("")

out.append("-- 今日の生産計画（動作確認用：500台）")
out.append("INSERT INTO production_plans (target_date, product_id, qty, note) VALUES")
out.append(f"  (CURDATE(), {PRODUCT_ID}, 500, '動作確認用');")
out.append("")

out.append("-- 在庫（動作確認用のデモ値。実際の数量は棚卸し画面で入力してください）")
demo_stock = [('バイオレット', 20000, 30), ('上白糖ES', 20000, 30), ('ケークドール', 20000, 30),
              ('エスイー57', 20000, 30), ('明治 業務用 酪農牛乳', 3000, 3),
              ('エースホイップV30', 3000, 3), ('お餅シートプレーンS-5.5-5.5', 2, 60)]
rows = [f"  ({mat_ids[n]}, {qty}, DATE_ADD(CURDATE(), INTERVAL {days} DAY), '冷蔵庫')"
        for n, qty, days in demo_stock if n in mat_ids]
out.append("INSERT INTO inventory (material_id, qty, expiry_date, location) VALUES")
out.append(",\n".join(rows) + ";")
out.append("")

out.append("-- 発注の実例（お餅シート／株式会社ホーライ宛）")
out.append("INSERT INTO purchase_orders (id, order_no, company_id, supplier_id, delivery_place, status,"
           " order_date, desired_date, note) VALUES")
out.append(f"  (1, 'PO-SAMPLE-001', 1, {sup_ids['株式会社ホーライ']}, '株式会社アルムハート', 'draft',"
           " CURDATE(), DATE_ADD(CURDATE(), INTERVAL 14 DAY), 'いただいた発注書の実例をもとにした確認用データ');")
out.append("INSERT INTO purchase_order_items (order_id, material_id, item_name, qty, unit, sort_no) VALUES")
out.append(f"  (1, {mat_ids['お餅シートプレーンS-5.5-5.5']}, 'お餅シートプレーンS-5.5-5.5', 1, 'ケース', 1);")
out.append("")

OUT.write_text("\n".join(out) + "\n", encoding='utf-8')
print('wrote', OUT)
print('材料', len(reg.materials), '業者', len(reg.suppliers), '部位', len(parts),
      '直接配合', len(direct))
